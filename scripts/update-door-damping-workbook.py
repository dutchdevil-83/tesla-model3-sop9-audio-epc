from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Tesla_Model_3_Highland_SOP9_Audio_BOM_v2.1.0.xlsx"
OUTPUT = ROOT / "Tesla_Model_3_Highland_SOP9_Audio_BOM_v2.2.0.xlsx"
DATA = ROOT / "data" / "door-damping" / "door_damping_v2.2.json"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
WHITE = "FFFFFF"


def clone_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)


def set_sheet_header(ws, title: str, subtitle: str, columns: list[str]) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="44546A")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 42

    for col, label in enumerate(columns, 1):
        cell = ws.cell(4, col, label)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}4"


def add_damping_bom(wb, data: dict) -> None:
    if "Door_Damping_BOM" in wb.sheetnames:
        del wb["Door_Damping_BOM"]
    ws = wb.create_sheet("Door_Damping_BOM", 4)
    columns = [
        "Scenario", "Layer", "Material", "Function / location", "Qty",
        "Bünde BOM nominal size", "Current verified size", "Thickness mm",
        "Area/sheet m²", "Total area m²", "Unit €", "Extended €",
        "Source", "Equivalence policy", "Engineering note", "URL"
    ]
    set_sheet_header(
        ws,
        "Door acoustic treatment BOM v2.2",
        "Front-only and all-four-door scenarios. Bünde kit quantities are the reference baseline; individual prices are sourcing observations and exclude shipping. Functional equivalence is controlled separately and is never thickness-only.",
        columns,
    )

    material_by_name = {m["material"]: m for m in data["materials"]}
    row = 5
    for scenario in data["scenarios"]:
        for name, qty in scenario["quantities"].items():
            if not qty:
                continue
            m = material_by_name[name]
            size = m.get("current_size", "")
            area = 0.0
            if "x" in size.lower() and name != "STP pressure roller":
                try:
                    nums = [float(x.strip()) for x in size.lower().replace("mm", "").split("x")[:2]]
                    area = nums[0] * nums[1] / 1_000_000
                except Exception:
                    area = 0.0
            values = [
                scenario["name"], m["layer"], name, m["function"], qty,
                m.get("bunde_size", ""), size, m.get("thickness_mm", 0), area,
                area * qty, m.get("unit_price_eur", ""),
                f"=IF(OR(E{row}=\"\",K{row}=\"\"),\"\",E{row}*K{row})",
                m.get("source", ""), m["equivalence"], m["gate"], m.get("url", "")
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            row += 1

        ws.cell(row, 1, scenario["name"])
        ws.cell(row, 2, "REFERENCE KIT")
        ws.cell(row, 3, "Bünde PRO" if scenario["name"] == "FRONT ONLY" else "Bünde XTREME")
        ws.cell(row, 12, scenario["reference_kit_eur"])
        ws.cell(row, 13, "CarHifi-Store Bünde")
        ws.cell(row, 14, "REFERENCE")
        ws.cell(row, 15, "Reference package price incl. VAT; shipping excluded.")
        ws.cell(row, 16, data["reference_kits"][0 if scenario["name"] == "FRONT ONLY" else 1]["url"])
        for c in ws[row]:
            c.fill = PatternFill("solid", fgColor=LIGHT)
            c.font = Font(bold=True)
        row += 2

    widths = [15, 24, 26, 52, 8, 30, 32, 12, 14, 14, 12, 14, 25, 24, 70, 55]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cells in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r in range(5, ws.max_row + 1):
        ws.cell(r, 8).number_format = "0.0"
        ws.cell(r, 9).number_format = "0.000"
        ws.cell(r, 10).number_format = "0.000"
        ws.cell(r, 11).number_format = '€ #,##0.00'
        ws.cell(r, 12).number_format = '€ #,##0.00'


def add_equivalence_sheet(wb, data: dict) -> None:
    if "Door_Damping_Equiv" in wb.sheetnames:
        del wb["Door_Damping_Equiv"]
    ws = wb.create_sheet("Door_Damping_Equiv", 5)
    columns = [
        "Baseline", "Candidate", "Market", "Observed price", "Package / coverage",
        "Technical evidence", "Decision", "Reason / remaining gate", "Seller URL",
        "Official / technical URL", "Checked"
    ]
    set_sheet_header(
        ws,
        "Door damping functional-equivalence matrix v2.2",
        "Candidates are judged by acoustic/mechanical function, material architecture, moisture/temperature suitability and published performance data. AliExpress hits are leads only until the exact variant has sufficient technical evidence.",
        columns,
    )
    row = 5
    for item in data["alternatives"]:
        price = item.get("price_eur")
        if price is None and item.get("price_usd") is not None:
            price = f"USD {item['price_usd']:.2f} (variant/checkout verify)"
        values = [
            item["baseline"], item["candidate"], item["market"], price,
            item.get("package", ""), item.get("technical", ""), item["decision"],
            item["reason"], item.get("url", ""), item.get("official", ""), data["checked"]
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        decision = item["decision"]
        fill = GREEN if "APPROVED" in decision and "NOT" not in decision else RED if "NOT APPROVED" in decision else AMBER
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row, 7).font = Font(bold=True)
        row += 1

    row += 1
    strict = [m for m in data["materials"] if m["equivalence"] == "STRICT"]
    ws.cell(row, 1, "STRICT MATERIAL GATES")
    ws.cell(row, 1).font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
    row += 1
    for m in strict:
        ws.cell(row, 1, m["material"])
        ws.cell(row, 6, m["gate"])
        ws.cell(row, 7, "NO GENERIC SUBSTITUTE APPROVED")
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=RED)
        ws.cell(row, 7).font = Font(bold=True)
        ws.cell(row, 9, m.get("url", ""))
        row += 1

    widths = [22, 36, 14, 23, 35, 58, 28, 70, 55, 55, 14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cells in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = f"A4:K{4 + len(data['alternatives'])}"


def add_product_master_rows(wb, data: dict) -> None:
    old = "EU_Price_Matrix_v2.1"
    new = "EU_Price_Matrix_v2.2"
    if old in wb.sheetnames:
        wb[old].title = new
    ws = wb[new]
    ws["A1"] = "EU price matrix v2.2 — NL / DE / BE / FR + door acoustic treatment"
    ws["A2"] = (
        "Snapshot 2026-09-10. Door-damping rows add verified EU observations and AliExpress research leads. "
        "AliExpress entries are not treated as country-low values until exact variant/technical equivalence is verified. "
        "Shipping, checkout, stock and warranty remain final-order checks."
    )

    rows = [
        ("B", "STP Aero Bomb", "Door damping / local CLD", None, 17.99, "PimpMySound", "https://pimpmysound.com/StP-Bomb-Aero-Matte", "HIGH", "STRICT baseline; no generic substitute approved."),
        ("B", "STP Black Gold 2.3 mm", "Door damping / distributed CLD", None, 14.90, "Carmedia-Shop", "https://www.carmedia-shop.de/STP-Black-Gold-Daemmung-1-Matte-0375-m", "HIGH", "Baseline: 100 µm foil; MLF 0.33; alternatives allowed by functional-equivalence gate."),
        ("B", "STP Accent 10", "Door damping / absorber", None, 14.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/stp-accent-10", "HIGH", "10 mm wet-zone-capable absorber/decoupler baseline."),
        ("B", "STP Sonora", "Door damping / multilayer barrier", None, 14.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/stp-sonora", "HIGH", "STRICT; multilayer moisture-resistant membrane/foam; plain foam not equivalent."),
        ("B", "STP Bromo", "Door damping / door-card barrier", None, 19.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/stp-bromo", "HIGH", "STRICT. Current standalone product is 7 mm; Bünde kit text still lists older 10 mm dimensions."),
        ("B", "Bünde Lautsprecher dämmsetPRO", "Door damping kit / front only", None, 199.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/lautsprecher-daemmsetpro", "HIGH", "Reference front-door package; shipping excluded."),
        ("B", "Bünde Lautsprecher dämmsetXTREME", "Door damping kit / all 4 doors", None, 299.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/lautsprecher-daemmsetxtreme", "HIGH", "Reference 2-4 door package including pressure roller; shipping excluded."),
        ("B", "CTK Premium 2.2 mm", "Door damping / CLD alternative", 6.90, None, "Toraz", "https://www.toraz.nl/dempmateriaal", "HIGH", "APPROVED - HIGH functional alternative to Black Gold: 2.2 mm, 100 µm foil, MLF 0.36, 3.07 kg/m²."),
        ("C", "Reckhorn ABX-tra 2.5 mm", "Door damping / CLD alternative", None, 34.99, "Reckhorn", "https://reckhorn.com/products/abx-tra-2-5mm-alubutyl", "MEDIUM", "CONDITIONAL: attractive €/m², but published MLF/foil thickness not captured."),
        ("C", "Reckhorn ABX 2.0 mm", "Door damping / CLD alternative", None, 29.99, "Reckhorn", "https://reckhorn.com/products/reckhorn-abx-2mm-alubutyl", "MEDIUM", "CONDITIONAL: economical but thinner and lacks published MLF/foil data."),
        ("B", "CTK SilenceFix 10 mm", "Door damping / absorber alternative", None, 109.00, "justSOUND", "https://just-sound.de/CTK", "HIGH", "APPROVED FUNCTIONAL - HIGH alternative to Accent 10; 12 sheets / 2.4 m²."),
        ("C", "AliExpress butyl/aluminium #1005012838442561", "Door damping research lead", None, None, "AliExpress / PriceArchive", "https://no.pricearchive.org/aliexpress.com/item/1005012838442561", "LOW", "NOT APPROVED. Indexed price USD 12.05; exact thickness/mass/foil/MLF variant not verified."),
        ("C", "AliExpress closed-cell foam #1005010659008580", "Door damping research lead", None, None, "AliExpress / PriceArchive", "https://ms.pricearchive.org/aliexpress.com/item/1005010659008580", "LOW", "CONDITIONAL research lead. Indexed price USD 10.93; 5/10 mm listing. Verify exact 10 mm variant, water/adhesive/temp/density data."),
        ("C", "STP pressure roller", "Installation tool", None, 15.00, "CarHifi-Store Bünde", "https://www.carhifi-store-buende.de/lautsprecher-daemmsetpro", "HIGH", "Generic hard pressure roller allowed."),
        ("C", "Door damping functional-equivalence dataset v2.2", "Engineering reference", None, None, "Project data", "data/door-damping/door_damping_v2.2.json", "HIGH", "Decision source for approved/conditional/strict damping substitutions."),
    ]

    ws.insert_rows(64, amount=len(rows))
    for idx, rowdata in enumerate(rows, 64):
        clone_row_style(ws, 63, idx, 30)
        priority, product, category, nl_price, de_price, seller, source, confidence, notes = rowdata
        values = [priority, product, category, "NO", None, None]
        for col, value in enumerate(values, 1):
            ws.cell(idx, col, value)
        if nl_price is not None:
            ws.cell(idx, 7, nl_price); ws.cell(idx, 8, seller); ws.cell(idx, 9, source)
            ws.cell(idx, 10, None); ws.cell(idx, 11, "NO CURRENT INDEXED LOCAL OFFER FOUND"); ws.cell(idx, 12, None)
        elif de_price is not None:
            ws.cell(idx, 7, None); ws.cell(idx, 8, "NO CURRENT INDEXED LOCAL OFFER FOUND"); ws.cell(idx, 9, None)
            ws.cell(idx, 10, de_price); ws.cell(idx, 11, seller); ws.cell(idx, 12, source)
        else:
            ws.cell(idx, 7, None); ws.cell(idx, 8, "NON-EU / RESEARCH LEAD"); ws.cell(idx, 9, source)
            ws.cell(idx, 10, None); ws.cell(idx, 11, "NO APPROVED EU PRICE MAPPED"); ws.cell(idx, 12, None)
        ws.cell(idx, 13, None); ws.cell(idx, 14, "NO CURRENT INDEXED LOCAL OFFER FOUND"); ws.cell(idx, 15, None)
        ws.cell(idx, 16, None); ws.cell(idx, 17, "NO CURRENT INDEXED LOCAL OFFER FOUND"); ws.cell(idx, 18, None)
        ws.cell(idx, 19, f'=IF(COUNT(G{idx},J{idx},M{idx},P{idx})=0,"",MIN(G{idx},J{idx},M{idx},P{idx}))')
        ws.cell(idx, 20, f'=IF(S{idx}="","",IF(G{idx}=S{idx},"NL",IF(J{idx}=S{idx},"DE",IF(M{idx}=S{idx},"BE","FR"))))')
        ws.cell(idx, 21, f'=IF(COUNT(G{idx},J{idx},M{idx},P{idx})=0,"",AVERAGE(G{idx},J{idx},M{idx},P{idx}))')
        ws.cell(idx, 22, f'=COUNT(G{idx},J{idx},M{idx},P{idx})')
        ws.cell(idx, 23, "")
        ws.cell(idx, 24, "")
        ws.cell(idx, 25, data["checked"])
        ws.cell(idx, 26, confidence)
        ws.cell(idx, 27, "DOOR DAMPING / VERIFY CHECKOUT")
        ws.cell(idx, 28, notes)
        ws.cell(idx, 29, source)
        ws.cell(idx, 30, "DOOR ACOUSTIC TREATMENT")

    ws.auto_filter.ref = f"A4:AD{63 + len(rows)}"


def update_overview(wb, data: dict) -> None:
    ws = wb["Overview"]
    ws["A1"] = "Tesla Model 3 Highland 2026 SOP9 - Audio Upgrade BOM v2.2"
    ws["A2"] = (
        "v2.2 adds optional door acoustic treatment with separate Front-only and All-4-doors BOMs, "
        "functional-equivalence gates, and researched EU/AliExpress alternatives. Existing SOP9 audio architecture, "
        "Tesla assets and MB Car Audio / HELIX comparison remain preserved."
    )
    ws["B9"] = data["checked"]
    ws["B10"] = 0.8613


def append_sources(wb, data: dict) -> None:
    ws = wb["Sources"]
    existing = {str(ws.cell(r, 3).value or "") for r in range(1, ws.max_row + 1)}
    urls = []
    for kit in data["reference_kits"]:
        urls.append((kit["name"], kit["url"], "Door damping reference kit / BOM"))
    for m in data["materials"]:
        urls.append((m["material"], m.get("url", ""), "Door damping baseline / technical sourcing"))
    for alt in data["alternatives"]:
        urls.append((alt["candidate"], alt.get("url", ""), "Door damping alternative / price observation"))
        if alt.get("official"):
            urls.append((alt["candidate"] + " official", alt["official"], "Manufacturer technical specification"))
    for name, url, note in urls:
        if not url or url in existing:
            continue
        r = ws.max_row + 1
        ws.cell(r, 1, name)
        ws.cell(r, 2, note)
        ws.cell(r, 3, url)
        existing.add(url)


def validate(wb) -> None:
    required = {"Door_Damping_BOM", "Door_Damping_Equiv", "EU_Price_Matrix_v2.2"}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise RuntimeError(f"Missing sheets: {sorted(missing)}")
    if wb["Door_Damping_BOM"]["C5"].value != "STP Aero Bomb":
        raise RuntimeError("Door_Damping_BOM baseline validation failed")
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        errors.append(f"{ws.title}!{cell.coordinate}")
    if errors:
        raise RuntimeError("Formula errors: " + ", ".join(errors[:10]))


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    data = json.loads(DATA.read_text(encoding="utf-8"))
    wb = load_workbook(SOURCE)
    update_overview(wb, data)
    add_product_master_rows(wb, data)
    add_damping_bom(wb, data)
    add_equivalence_sheet(wb, data)
    append_sources(wb, data)
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    validate(wb)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT.name} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    main()
